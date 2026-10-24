#!/usr/bin/env python3
"""Offline meeting media from saved observations; never imports robot/model code."""
from pathlib import Path
import argparse
import csv
import hashlib
import json
import math
import subprocess

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

REPO = Path(__file__).resolve().parents[1]
REPORTS = REPO / "reports"
FONT = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
W, H = 1920, 1080
SOURCES = {}


def track(path):
    path = Path(path).resolve()
    SOURCES[str(path)] = hashlib.sha256(path.read_bytes()).hexdigest()
    return path


def read(path):
    return json.loads(track(path).read_text())


def font(size):
    return ImageFont.truetype(FONT, size)


def wrapped(draw, text, xy, width, size=30, fill="#cbd5e1", max_lines=8):
    x, y = xy
    lines = []
    for paragraph in str(text).splitlines():
        line = ""
        for ch in paragraph:
            if draw.textlength(line + ch, font=font(size)) > width:
                lines.append(line)
                line = ch
            else:
                line += ch
        lines.append(line)
    if len(lines) > max_lines:
        raise ValueError(f"Text exceeds slide area: {text}")
    for line in lines:
        draw.text((x, y), line, font=font(size), fill=fill)
        y += size + 12
    return y


def add_picture(canvas, path, box, plan=None):
    if path is None:
        d = ImageDraw.Draw(canvas)
        wrapped(d, "没有动作后 observation\n未用其他帧替代", (box[0]+45, box[1]+160),
                box[2]-90, 36)
        return
    with Image.open(track(path)) as src:
        im = src.convert("RGB")
    if plan and plan.get("decision") == "ACTION":
        a, b = plan["grasp_px"], plan["target_px"]
        d = ImageDraw.Draw(im)
        d.line([tuple(a), tuple(b)], fill="#fbbf24", width=5)
        angle = math.atan2(b[1]-a[1], b[0]-a[0])
        wings = [(b[0]-20*math.cos(angle+t), b[1]-20*math.sin(angle+t)) for t in (-.5, .5)]
        d.polygon([tuple(b), *wings], fill="#fbbf24")
        for point, label, color in ((a, "Pa", "#fb7185"), (b, "Pb", "#22d3ee")):
            x, y = point
            d.ellipse((x-7, y-7, x+7, y+7), fill=color, outline="white", width=2)
            d.text((x+9, y+3), label, font=font(23), fill=color, stroke_width=2, stroke_fill="black")
    im.thumbnail((box[2], box[3]), Image.Resampling.LANCZOS)
    canvas.paste(im, (box[0]+(box[2]-im.width)//2, box[1]+(box[3]-im.height)//2))


class Deck:
    def __init__(self, out, key, title, selection, sources):
        self.out, self.key, self.title = out, key, title
        self.folder = out/"slides"/key
        self.folder.mkdir(parents=True)
        self.meta = dict(id=key, title=title, selection_reason=selection,
                         source_runs=sources, slides=[], duration_s=0)

    def slide(self, subtitle, notes, pictures=(), duration=6, evidence=(), labels=None, plan=None):
        im = Image.new("RGB", (W, H), "#101827")
        d = ImageDraw.Draw(im)
        d.rectangle((0, 0, W, 12), fill="#38bdf8")
        wrapped(d, self.title, (55, 34), 1810, 44, "white", 1)
        wrapped(d, subtitle, (55, 110), 1810, 30, "#7dd3fc", 2)
        labels = labels or []
        if len(pictures) == 2:
            for i, pic in enumerate(pictures):
                x = 40+i*950
                wrapped(d, labels[i] if labels else ("BEFORE" if i == 0 else "AFTER"),
                        (x+15, 193), 880, 29, "white", 1)
                add_picture(im, pic, (x, 235, 930, 525), plan if i == 0 else None)
            wrapped(d, notes, (55, 790), 1810, 29, max_lines=5)
        elif len(pictures) == 1:
            add_picture(im, pictures[0], (40, 200, 1250, 760))
            wrapped(d, notes, (1320, 218), 550, 29, max_lines=16)
        else:
            wrapped(d, notes, (85, 250), 1730, 40, max_lines=12)
        d.line((55, 1010, 1865, 1010), fill="#334155", width=2)
        wrapped(d, "2026-10-20  |  保存的 OBS / 结果图组成的幻灯片；非连续运动录像，非真实耗时回放",
                (55, 1025), 1730, 23, "#94a3b8", 1)
        number = len(self.meta["slides"])+1
        path = self.folder/f"{number:03d}.png"
        im.save(path)
        self.meta["slides"].append(dict(index=number, title=subtitle,
            file=str(path.relative_to(self.out)), duration_s=duration,
            pictures=[str(Path(p).resolve()) if p else None for p in pictures],
            evidence=[str(track(p)) for p in evidence], notes=notes))
        self.meta["duration_s"] += duration


def episode(out, key, title, names, selection, conclusion, highlight=None):
    roots = [REPORTS/n for n in names]
    deck = Deck(out, key, title, selection, names)
    summaries = [read(p/"episode_summary.json") for p in roots]
    deck.slide("真实机器人实验 · 代表 run 的选择依据",
        selection+"\n\n"+conclusion+"\n\n所有抓取 acquisition / 最终 FOLD_READY 均未在此材料中自动确认。",
        evidence=[p/"episode_summary.json" for p in roots], duration=7)
    init = roots[0]/"initialization"
    if (init/"before/rgb.png").exists() and (init/"after/rgb.png").exists():
        deck.slide("初始化：抓起再丢下", "初始化 before / after；执行完成不等于已确认抓取成功。",
            [init/"before/rgb.png", init/"after/rgb.png"], labels=["初始化前", "初始化后"])
    first = None
    last = None
    for seg, root in enumerate(roots):
        if seg:
            deck.slide(f"人工授权续跑 · 第 {seg+1} 段",
                "上一段已经停止，本页之后是用户明确授权的新段。\n未重新初始化；重新采集 observation 后重新规划。\n"
                +("本次 goal-only 续跑采用右移 20 mm 的观测位姿。" if "goal_only" in key else
                  "保留原安全检查；不是未中断的自主连续运行。"),
                duration=5)
        for folder in sorted(root.glob("round_[0-9][0-9][0-9]")):
            before, after = folder/"before/rgb.png", folder/"after/rgb.png"
            if not before.exists():
                continue
            first = first or before
            last = after if after.exists() else before
            plan = read(folder/"plan.json") if (folder/"plan.json").exists() else {}
            execution = read(folder/"execution.json") if (folder/"execution.json").exists() else {}
            timing = read(folder/"timing.json") if (folder/"timing.json").exists() else {}
            status = execution.get("status", "NO_SAVED_PLAN")
            decision = plan.get("decision", "未调用模型")
            latency = timing.get("planning_s", timing.get("agent_wall_s"))
            timing_text = f"规划 {latency:.2f}s" if latency is not None else "规划耗时未记录"
            note = f"{decision} / {status}；{timing_text}；机器人 {timing.get('robot_execution_s', 0):.2f}s。"
            reason = plan.get("short_reason", plan.get("short_rationale", "本轮没有保存的模型计划。"))
            # Long rationale remains available in plan.json; avoid illegible slide text.
            if len(reason) > 125:
                reason = reason[:122]+"…"
            note += "\n模型原理由："+reason
            if execution.get("error"):
                note += "\n拒绝/中止："+execution["error"][:150]
            elif plan.get("supporting_experiences"):
                note += f"\n引用历史经验 {len(plan['supporting_experiences'])} 条；完整引用见原 plan.json。"
            if highlight == (root.name, folder.name):
                note += "\n复盘重点：用户指出已展开区域被覆盖；前后图可见衣身重新聚拢。布层运动机制仍需人工判断。"
            deck.slide(f"第 {seg+1} 段 / {folder.name}  ·  {status}", note,
                [before, after if after.exists() else None], duration=9 if highlight == (root.name, folder.name) else 6,
                evidence=[p for p in [folder/"plan.json", folder/"execution.json", folder/"timing.json"] if p.exists()],
                labels=["动作前（箭头只表示提案）", "实际动作后" if after.exists() else "未采集动作后图"],
                plan=plan)
    if first:
        actions = sum(s["physical_actions"] for s in summaries)
        deck.slide("初始与最终 observation", conclusion+f"\n本视频包含 {actions} 次展开动作；初始化另外计数。"
            "\n不同方法起始衣物状态不同；单次 run 不能作为成功率比较。",
            [first, last], labels=["初始化后的首个 observation", "最后保存的 observation"], duration=8)
    return deck


def offline(out):
    decks = []
    p = REPORTS/"world_model_probe_v1"
    d = Deck(out, "06_latent_dynamics", "离线 latent dynamics 诊断", "使用唯一完整的固定协议 run，未选择最优 seed。",
             [p.name])
    d.slide("132 transitions / 15 sessions / 11 groups", "Frozen DINOv2-small CLS + 小 MLP，5 seeds。\n"
        "按完整 group 划分；只训练此实验的小网络。\nAction-conditioned 未稳定优于 State-only：UNRESOLVED。\n"
        "本视频展示历史数据与离线结果，没有本实验生成的机器人动作。", evidence=[p/"summary.md"])
    rows = list(csv.DictReader(track(p/"transitions.csv").open()))
    row = next(r for r in rows if r["split"] == "test")
    d.slide("数据示例：CSV 中首个 held-out transition（非按效果挑选）",
        f"来源：{row['id']}\n真实历史 before/after；这里只说明数据形式，不能代表模型预测 RGB。",
        [Path(row["before_rgb"]), Path(row["after_rgb"])], labels=["历史 BEFORE", "历史 AFTER"])
    for name, subtitle, note in [
        ("baseline_comparison.png", "三个 baseline", "Test cosine：\nIdentity 0.88857\nState-only 0.91268\nAction-conditioned 0.90545\n\nC 优于 Identity，但低于 B。"),
        ("real_vs_shuffled.png", "Action sensitivity", "真实 action 优于 shuffled。\n但 C 仍未超过 state-only。\n\n使用了 action ≠ 学到更好的预测。\n不是物理 counterfactual。"),
        ("learning_curve.png", "按完整 group 的 learning curve", "固定 held-out split。\n80% → 100% 未证实继续改善。\n\n不能仅凭平台归因于 representation。")]:
        d.slide(subtitle, note, [p/name], evidence=[p/"summary.md"])
    decks.append(d)
    p = REPORTS/"trajectory_memory_agent_probe"
    d = Deck(out, "07_memory_no_memory_probe", "Memory / no-memory 单步建议对照", "全部 4 个新 query，避免按动作好坏筛选。",
             [p.name])
    d.slide("4 query；两条件独立会话；计划未执行", "Memory：139 条经验，平均 107.8s。\nNo-memory：平均 22.4s。\n"
        "人工动作质量比较待完成；读取经验不等于经验造成了改善。", evidence=[p/"report.md"])
    for q in sorted((p/"queries").glob("q*")):
        d.slide(q.name+" · 同一 observation 的两种提案", "仅离线建议，两种动作均未执行；原始 query 一致。",
            [q/"memory/proposed_action_overlay.png", q/"no_memory/proposed_action_overlay.png"],
            labels=["Memory", "No-memory"],
            evidence=[q/"memory/plan.json", q/"no_memory/plan.json"])
    decks.append(d)
    p = REPORTS/"teacher_experience_retrieval_probe_v1"
    d = Deck(out, "08_teacher_retrieval", "视觉检索能否找回 teacher 引用", "正例排名最好和最差的 query，加完整汇总。",
             [p.name])
    d.slide("7 queries / 139 references / 12 citations", "Frozen DINOv2-small，仅 historical BEFORE RGB。\n"
        "Hit@1/3/5 = 0；Hit@10 = 28.6%；Hit@20 = 57.1%。\n"
        "Best teacher rank 中位数 13，最差 122。\n"
        "检索 median 27.17 ms；WEAK_RETRIEVAL_SIGNAL。\n不是 state matching 或动作正确性评分。", evidence=[p/"summary.md"])
    for name, title, note in [
        ("hit_and_citation_recall.png", "固定 K 的汇总", "没有根据结果调 K 或 encoder。\n7 个 query 来自同一 episode。"),
        ("contact_sheets/round_004_top20.png", "最好案例：round_004，best rank = 8", "即使最好案例，也没有进入 Top-5。\nTeacher citation 不等于最优动作标签。"),
        ("contact_sheets/round_005_top20.png", "最差案例：round_005，best rank = 122", "Teacher 引用未出现在 Top-20。\n这里忠实保留原视觉排序。")]:
        d.slide(title, note, [p/name], duration=8)
    decks.append(d)
    p = REPORTS/"one_shot_goal_planner_probe_v1"
    d = Deck(out, "09_one_shot_offline_probe", "One-shot 与原 goal-only Agent 对照", "展示相近抓点、最大坐标差、DONE 三类，完整十轮保留在原报告。",
             [p.name])
    d.slide("10 个保存的 observation；每 query 只调用一次", "One-shot median 21.57s / mean 21.35s / p95 37.87s。\n"
        "原 Agent median 61.97s；均值之比加速 2.69×。\n"
        "Decision agreement 100%，但人工质量评审 0/10。\n"
        "输入信息也减少，不能把所有差异只归因于工具循环。", evidence=[p/"summary.md"])
    d.slide("逐 query 的 wall latency", "原 Agent 使用保存耗时。\n未重新运行原 Agent。\n未执行 one-shot 建议。",
            [p/"latency_comparison.png"])
    for qid, label in [("initial_round_000", "抓点接近的案例"),
                       ("continue_round_005", "Pa/Pb 差异最大的案例"),
                       ("continue_round_006", "双方均 DONE")]:
        q=p/"queries"/qid
        d.slide(qid+" · "+label, "原 plan 不是 ground truth；箭头差异只作描述，质量需要人工比较。",
            [q/"original_overlay.png", q/"one_shot_overlay.png"],
            labels=["原 goal-only Agent", "One-shot"], evidence=[q/"original_plan.json", q/"one_shot_plan.json"])
    decks.append(d)
    p=REPORTS/"host_grasp_geometry_audit_v2"
    d=Deck(out,"10_host_geometry_audit","Host 抓点几何离线审计","全部12轮复现；重点展示缺深度、坡面与已执行对照。",[p.name])
    d.slide("12/12 runtime gate 复现一致；未部署新规则", "8 个 EXECUTED 对照，3 个 depth reject，1 个 IK reject。\n"
        "EXECUTED 只表示允许执行，抓取成功仍 UNKNOWN。\n"
        "round 9/10：局部坡面解释较合理；round 7：邻近像素可通过原 depth gate。\n"
        "不把 depth-valid 当成完整动作已获执行认证。", evidence=[p/"summary.md"])
    for name,title,note in [
        ("round_07/snap_detail.png","Round 7：中心缺深度","最近候选 [417,202]，偏移 1.414px。\n原点缺深度，真实 world XY 位移未知。\n只做离线候选搜索。"),
        ("round_09/diagnostic.png","Round 9：raw range 8.099mm","去坡度 residual RMSE 0.311mm。\n5×5 更像连续斜面。\n没有新增 residual 阈值。"),
        ("round_10/diagnostic.png","Round 10：raw range 8.109mm","去坡度 residual RMSE 0.331mm。\n更大窗口仍有曲率。\n不能推广为所有断层可放行。"),
        ("reference_distribution.png","已 EXECUTED 样本分布","缺少确认的真断层负例。\n不能据此宣称新分类器准确。\nIK case 单独保留。")]:
        d.slide(title,note,[p/name],duration=8)
    decks.append(d)
    return decks


def encode(out, decks):
    for deck in decks:
        meta=deck.meta
        concat=out/"slides"/deck.key/"frames.ffconcat"
        lines=["ffconcat version 1.0"]
        for s in meta["slides"]:
            lines += [f"file '{Path(s['file']).name}'",f"duration {s['duration_s']}"]
        lines += [f"file '{Path(meta['slides'][-1]['file']).name}'"]
        concat.write_text("\n".join(lines)+"\n")
        target=out/"videos"/f"{deck.key}.mp4"
        subprocess.run(["ffmpeg","-hide_banner","-loglevel","error","-f","concat","-safe","0","-i",str(concat),
            "-vf","fps=15,tpad=stop_mode=clone:stop_duration=3",
            "-frames:v",str(meta["duration_s"]*15),"-c:v","libx264","-preset","fast","-crf","20",
            "-threads","2","-pix_fmt","yuv420p","-movflags","+faststart",str(target)],check=True)
        probe=json.loads(subprocess.check_output(["ffprobe","-v","error","-show_entries",
            "stream=codec_name,width,height,pix_fmt,nb_frames:format=duration","-of","json",str(target)]))
        assert abs(float(probe["format"]["duration"])-meta["duration_s"])<.1
        meta.update(video=str(target.relative_to(out)),video_probe=probe)
        print("ENCODED",deck.key,meta["duration_s"],flush=True)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output",type=Path,default=REPORTS/"meeting_materials_20261020")
    args=parser.parse_args()
    out=args.output.resolve();out.mkdir(parents=True,exist_ok=False)
    (out/"videos").mkdir()
    decks=[
        episode(out,"01_full_memory","历史 trajectory memory · 真实闭环",
            ["trajectory_memory_full_episode_v1_collection_aligned"],
            "唯一完成多步执行的 collection-aligned memory run；139 条冻结经验。",
            "7 次展开动作后 USER_STOP；平均 Agent 规划 120.89s。人工 FOLD_READY 未判定。"),
        episode(out,"02_goal_only_agent","Goal-only Agent · 原段与授权续跑",
            ["goal_only_full_episode_v1_restart02","goal_only_full_episode_v1_continue01"],
            "展示从新初始化到模型 DONE 的完整记录，包括拒绝和人工恢复。",
            "2+6 次展开动作；10 次规划。先 SAFETY_STOP，授权换视角续跑后 DONE；非全程无干预。"),
        episode(out,"03_one_shot_12round","One-shot policy · 完整 12 轮与安全停止",
            ["one_shot_goal_full_episode_v1_restart01"]+
            [f"one_shot_goal_full_episode_v1_continue0{i}" for i in range(1,5)],
            "保留全部12轮、4次拒绝和用户授权续跑，不筛去失败动作。",
            "累计 8 次展开动作、4 次拒绝，最终 ROUND_LIMIT；平均规划 21.74s。人工质量待判。"),
        episode(out,"04_one_shot_regression","One-shot policy · 布层覆盖退化案例",
            ["one_shot_goal_full_episode_v1_restart02"],
            "本次讨论的代表失败 run：round_003 前后出现明显重新聚拢。",
            "6 次展开动作；第7轮159.28mm提案超过150mm而 SAFETY_STOP。不能自动判定成功。",
            highlight=("one_shot_goal_full_episode_v1_restart02","round_003")),
        episode(out,"05_one_shot_depth_stop","One-shot policy · 最新深度拒绝案例",
            ["one_shot_goal_full_episode_v1_20261020_220217_032210193"],
            "用户独立终端运行的最新记录；保留初始化、1次执行和未执行提案。",
            "round_001：Pa=[551,230]，5×5 高度差11.046mm > 8mm。原始 gate 不能区分坡面和断层。")
    ]
    decks += offline(out)
    prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
    for deck in decks:
        for s in deck.meta["slides"]:
            slide=prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_picture(str(out/s["file"]),0,0,width=prs.slide_width,height=prs.slide_height)
            slide.notes_slide.notes_text_frame.text=s["title"]+"\n"+s["notes"]+"\nSources:\n"+"\n".join(s["evidence"]+[
                x for x in s["pictures"] if x])
    prs.save(out/"all_observation_slides.pptx")
    encode(out,decks)
    (out/"source_hashes.json").write_text(json.dumps(SOURCES,ensure_ascii=False,indent=2))
    manifest=dict(date="2026-10-20",format="observation slideshow, not continuous video",
        resolution=[W,H],fps=15,model_calls=0,robot_connections=0,
        decks=[d.meta for d in decks],source_files_unchanged=all(
            hashlib.sha256(Path(p).read_bytes()).hexdigest()==h for p,h in SOURCES.items()))
    assert manifest["source_files_unchanged"]
    (out/"manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
    print("COMPLETE",out,flush=True)


if __name__=="__main__":
    main()
