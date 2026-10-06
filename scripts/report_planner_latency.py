"""Produce the evidence-backed Chinese PPTX, tables and review for latency ablations."""
from __future__ import annotations
import argparse, csv, json, math, sys
from pathlib import Path
from collections import Counter
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from cloth_agent.harness.planner_profile import analyze_events

LABELS={'A_files_split':'A 文件上下文·两次调用','B_inline_split':'B 内联上下文·两次调用',
        'C_inline_merged':'C 内联上下文·合并调用','D_ordered':'D 明确决策步骤','E_geometry':'E 步骤 + 几何代码',
        'F_concise':'F 简洁输出','G_direct':'G 图片直接输入','C_repeat':'C′ 合并调用复测',
        'H_low_effort':'H 同模型·low effort','I_matched_inline':'I 实际已读内容内联'}
BG='F4F6FA'; INK='182D47'; MUTED='63758B'; BLUE='2563EB'; TEAL='008577'; ORANGE='D07825'; WHITE='FFFFFF'

def read(p):return json.loads(Path(p).read_text())
def write(p,d):Path(p).write_text(json.dumps(d,ensure_ascii=False,indent=2),encoding='utf-8')
def code(row):return 'C′' if row['arm']=='C_repeat' else row['arm'].split('_')[0]
def inside(p,poly):
    x,y=p;result=False
    for i,a in enumerate(poly):
        b=poly[i-1]
        if (a[1]>y)!=(b[1]>y) and x<(b[0]-a[0])*(y-a[1])/(b[1]-a[1])+a[0]:result=not result
    return result

def aggregate(out):
    reports=read(out/'results.json');rubric=read(out/'quality_rubric.json');rows=[]
    for r in reports:
        calls=r['calls'];counts=Counter();rounds=0;wait=emit=cli=0.;tokens=0;all_models_tokens=0;output=0;thinking=0;cost=0.;models=set();retries=[];usage_caveats=[]
        for c in calls:
            counts.update(c['tool_counts']);u=c.get('usage') or {}
            tokens+=sum(u.get(k,0) or 0 for k in ['input_tokens','output_tokens','cache_read_input_tokens','cache_creation_input_tokens'])
            output+=u.get('output_tokens',0) or 0;cost+=c.get('cost_usd') or 0
            thinking+=(u.get('output_tokens_details') or {}).get('thinking_tokens',0) or 0
            models.update((c.get('models') or {}).keys())
            all_models_tokens+=sum(sum(mu.get(k,0) or 0 for k in ['inputTokens','outputTokens','cacheReadInputTokens','cacheCreationInputTokens']) for mu in (c.get('models') or {}).values())
            trace=Path(c['trace']);raw=(trace/'claude_stdout.txt').read_text();events=analyze_events(raw)
            for line in raw.splitlines():
                event=json.loads(line)
                if event.get('type')=='system' and event.get('subtype')=='api_retry':
                    retries.append({'stage':c['stage'],'attempt':event.get('attempt'),'status':event.get('error_status'),
                                    'delay_ms':event.get('retry_delay_ms'),'elapsed_s':event.get('_cloth_timing',{}).get('elapsed_s')})
                if event.get('type')=='stream_event' and event.get('event',{}).get('type')=='message_delta':
                    uevent=event['event'].get('usage',{})
                    if 'iterations' in uevent and uevent['iterations']==[] and uevent.get('output_tokens',999)>0 and uevent.get('output_tokens',999)<=5:
                        usage_caveats.append({'stage':c['stage'],'output_tokens':uevent['output_tokens'],
                            'note':'Very small output count with empty iterations; may be incomplete usage. Do not infer reasoning work from this count.'})
            rounds+=len(events['rounds']);wait+=sum(x['wait_before_first_event_s'] for x in events['rounds'])
            emit+=sum(x['emission_s'] or 0 for x in events['rounds'])
            timing=read(trace/'timing.json');cli+=timing['phases_s'].get('remote_claude',0)
            write(trace.parent/'public_timing.json',events)
        val=r.get('validation',{});gp=val.get('grasp_pixel_xy');tg=val.get('target_mappings',[])
        # Original upright image_0 -> prepared collar-up image_6.
        grip=[1279-gp[1],gp[0]] if gp else None
        target=[1279-tg[0]['grounding_pixel_xy'][1],tg[0]['grounding_pixel_xy'][0]] if tg else None
        q={'grasp_on_requested_sleeve':inside(grip,rubric['grasp_sleeve_polygon']) if grip else None,
           'grasp_in_preferred_cuff_band':grip[0]>=rubric['preferred_cuff_band_x_min'] if grip else None,
           'target_on_chest':inside(target,rubric['target_chest_polygon']) if target else None,
           'inward_transport':target[0]<grip[0] if grip and target else None}
        semantic=all(v is True for v in q.values())
        row={'arm':r['arm'],'label':LABELS[r['arm']],'status':r['status'],'wall_s':r['wall_s'],
             'calls':len(calls),'rounds':rounds,'read_calls':counts['Read'],
             'view_calls':counts['mcp__cloth_image__view_image'],'edits':sum(v for k,v in counts.items() if k.endswith(('crop_image','rotate_image','resize_image'))),
             'response_wait_s':wait,'emission_window_s':emit,'remote_cli_s':cli,'tokens':tokens,'all_models_reported_tokens':all_models_tokens,'output_tokens':output,'reported_thinking_tokens':thinking,
             'reported_cost_usd':cost,'models':sorted(models),'grasp_id':val.get('eligible_reference'),
             'grasp_collar_up':grip,'target_collar_up':target,'offline_quality_checks':q,
             'quality':'离线检查通过' if r['status']=='COMPLETED' and semantic else '需复核', 'error':r.get('error'),
             'api_retry_count':len(retries),'api_retries':retries,'token_usage_caveats':usage_caveats}
        rows.append(row)
    write(out/'analysis.json',rows)
    return rows


class Deck:
    def __init__(self):
        self.prs=Presentation();self.prs.slide_width=Inches(13.333);self.prs.slide_height=Inches(7.5)
    def text(self,s,text,x,y,w,h,size=20,color=INK,bold=False):
        shape=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=shape.text_frame
        tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
        for i,line in enumerate(str(text).split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.space_after=Pt(9)
            p.font.name='Noto Sans CJK SC';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        return shape
    def rect(self,s,x,y,w,h,color):
        shape=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(color);shape.line.fill.background();return shape
    def slide(self,title,subtitle='',source=''):
        s=self.prs.slides.add_slide(self.prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
        self.rect(s,0,0,13.333,.12,BLUE);self.text(s,title,.55,.35,12.2,.55,29,bold=True)
        self.text(s,subtitle,.58,1.04,12,.6,14,MUTED)
        self.text(s,'Cloth Agent  ·  Latency Study  ·  2026-10-05',.58,7.1,8,.22,10,MUTED)
        self.text(s,str(len(self.prs.slides)),12.05,7.07,.5,.25,11,MUTED)
        if source:s.notes_slide.notes_text_frame.text='数据来源：'+source+'\n所有实验离线重放，不执行机器人。原始真实轮次另列。'
        return s
    def table(self,s,headers,rows,widths=None,y=1.78,h=4.75,size=16):
        sh=s.shapes.add_table(len(rows)+1,len(headers),Inches(.58),Inches(y),Inches(12.15),Inches(h))
        t=sh.table
        if widths:
            for c,w in zip(t.columns,widths):c.width=Inches(w)
        for i,row in enumerate([headers]+rows):
            for j,val in enumerate(row):
                c=t.cell(i,j);c.text=str(val);c.margin_left=Inches(.12);c.margin_right=Inches(.08);c.margin_top=Inches(.05);c.margin_bottom=Inches(.02)
                c.fill.solid();c.fill.fore_color.rgb=RGBColor.from_string(INK if i==0 else WHITE if i%2 else 'E9EEF6')
                for p in c.text_frame.paragraphs:
                    p.font.name='Noto Sans CJK SC';p.font.size=Pt(size);p.font.bold=i==0;p.font.color.rgb=RGBColor.from_string(WHITE if i==0 else INK)
        return t
    def bullets(self,s,items,y=1.85,size=22):
        for i,item in enumerate(items):
            self.rect(s,.62,y+i*1.12,.08,.45,BLUE);self.text(s,item,.88,y+i*1.12,11.5,1.02,size)


def build(out,original):
    rows=aggregate(out);by={r['arm']:r for r in rows};deck=Deck()
    spans=read(original/'summary.json')['spans'];lookup={s['id']:s for s in spans};total=read(original/'summary.json')['elapsed_s']
    groups=[('初始拍照与感知',7),('操作前 supervisor',22),('朝向准备与 Molmo',29),('完整规划',36),('机器人执行及录像收尾',76),('操作后拍照',125),('视频证据准备',140),('evaluation',141),('操作后 supervisor',153),('经验更新与保存',163)]
    times=[(name,lookup[i]['duration_s']) for name,i in groups];times.append(('其他校验、调度、落盘',total-sum(v for _,v in times)))
    valid=[r for r in rows if r['quality']=='离线检查通过'];best=min(valid,key=lambda x:x['wall_s']) if valid else min(rows,key=lambda x:x['wall_s'])
    base=by['A_files_split'];c=by['C_inline_merged'];cr=by['C_repeat']
    control='C_repeat' if c['api_retry_count'] and not cr['api_retry_count'] else 'C_inline_merged'
    s=deck.slide('耗时来自结构、方法，还是 Agent？','真实轮次完整计时 + 固定证据离线对照  |  可复查的实验报告',str(out/'analysis.json'))
    deck.text(s,'25 分 22 秒',.65,1.95,7,.8,46,BLUE,True)
    deck.text(s,'原始真实轮次总耗时；其中规划 735.2 秒',.7,2.92,11,.65,25)
    deck.text(s,f'{len(rows)} 组对照  /  {sum(r["calls"] for r in rows)} 次完成记录的模型调用',.7,4.0,11,.65,28,TEAL,True)
    deck.text(s,'固定同一现场、同一组图片、同一模型；无机器人动作。\n本报告区分观测结果、因果线索与尚未验证的结论。',.7,5.1,11,1.2,22)
    s=deck.slide('先看实验结论','最快结果必须同时通过离线质量检查；不等于真实抓取成功',str(out/'analysis.json'))
    deck.bullets(s,[f'最快通过离线检查：{best["label"]}，{best["wall_s"]:.1f} 秒。',
                   f'基准 A：{base["wall_s"]:.1f} 秒；合并 C：{c["wall_s"]:.1f} 秒（API 重试 {c["api_retry_count"]} 次）。',
                   f'同配置复测 C′：{cr["wall_s"]:.1f} 秒；服务与执行波动必须计入解释。',
                   '结论限定为一个场景的筛查；不能宣称算法下限或物理精度不变。'])
    s=deck.slide('原始真实轮次：完整耗时表','各行互不重叠，可相加。1521.7 秒 = 25 分 22 秒。',str(original/'host_spans.csv'))
    deck.table(s,['阶段','耗时 / 秒','占比'],[[n,f'{v:.1f}',f'{100*v/total:.1f}%'] for n,v in times]+[['合计',f'{total:.1f}','100%']],widths=[8.05,2.05,2.05],h=5.12,size=15)
    s=deck.slide('何时算“这一轮结束”？','同一轮可以有不同的业务终点，不能混用时间口径。',str(original/'host_spans.csv'))
    deck.table(s,['终点','自开始累计','含义'],[['规划完成','约 941 秒','抓取点、目标点和动作已生成'],['机器人执行结束','约 1040 秒','包含执行与录像收尾'],['evaluation 完成','约 1216 秒','20 分 16 秒，已有本轮效果判断'],['整轮正常返回','约 1522 秒','又做了 supervisor 与经验更新']],widths=[3.2,2.35,6.6],h=3.2,size=19)
    deck.text(s,'evaluation 之后额外约 305 秒，不能归因于“选点太慢”。',.7,5.48,12,1,25,BLUE,True)
    s=deck.slide('735 秒规划：编辑结束后仍有 502 秒','阶段边界由 Host 计时；不是对私有思考内容的猜测。',str(original/'host_spans.csv'))
    deck.table(s,['子阶段','耗时','实际交付'],[['图像编辑与证据准备','231.8 秒','图片、发现、充分性判断、引用链'],['只读抓取点选择','261.4 秒','Rxx 抓取点与策略'],['只读目标及动作规划','241.1 秒','目标像素、动作序列、编译'],['其他规划开销','0.9 秒','调度与准备']],widths=[4.05,2.05,6.05],h=3.4,size=19)
    deck.text(s,'后两次调用均为 0 次编辑；但分别读取上下文 16 次。',.7,5.65,12,.85,23,TEAL,True)
    s=deck.slide('这次发现的混杂因素：可用信息 ≠ 实际已读信息','因此增加 I：只内联 A 的实际 Read 返回内容，继续保留两次规划调用。',str(out/'methodology.md'))
    deck.table(s,['组别','实际输入变化','观察与解释'],[['A 文件读取','自主选择；没有读 experience_context','基准的实际信息少于全部可用文件'],['B 全部内联','把此前未读的经验也交给模型','更快，但目标点明显左移，需复核'],['I 已读内容内联','只预先提供 A 实际读到的文本','更接近隔离上下文读取的交互成本']],widths=[2.7,4.5,4.95],h=3.3,size=18)
    deck.text(s,'B 的公开输出引用经验要求扩大搬运距离；这提供解释线索，尚不能单独证明因果。',.7,5.6,12,1,21,ORANGE)
    s=deck.slide('原始图像编辑：8 轮交互，3 次编辑','响应输出窗口包括 thinking、文字与工具参数；并非图片运算耗时。',str(original/'claude_calls.json'))
    edit=[['1','目录与图片列表','34.3','1.3'],['2','读取约束、参考状态等','1.8','3.1'],['3','查看原图和标记图','1.7','3.3'],['4','查看两张参考图','8.0','29.1'],['5','原图旋转 90°','8.8','13.9'],['6','裁剪袖子与已折区域','11.4','6.3'],['7','读取候选、查看 Rxx 图','11.0','13.8'],['8','提交 READY 与证据','12.0','40.6']]
    deck.table(s,['轮次','公开操作','响应前等待 / 秒','输出窗口 / 秒'],edit,widths=[.8,6.05,2.65,2.65],h=4.9,size=16)
    s=deck.slide('三类原因，分别怎样验证','速度收益只有在任务和质量约束可比时才有意义。',str(out/'design.json'))
    deck.table(s,['类别','控制变量','主要比较'],[['结构','任务和内容不变，改变交付与调用组织','A → B；B → C'],['方法 / 算法','结构固定，改变决策步骤或计算辅助','C → D；D → E'],['Agent 执行方式','方法固定，约束输出或消除取图交互','C → F；C → G'],['波动控制','同一配置再次运行','C → C′']],widths=[2.0,7.0,3.15],h=3.9,size=18)
    deck.text(s,'另测同模型 low effort；未测不同模型或 provider 内部排队。',.65,6.15,12,.5,17,ORANGE)
    s=deck.slide('可比性与质量约束','同一组 9 张图片，按校验值冻结；没有提供上次最终选点答案。',str(out/'input_hashes.json')+'; '+str(out/'quality_rubric.json'))
    deck.bullets(s,['固定模型请求、图片、任务、schema；所有调用禁止图像编辑。',
                   '前两组 motion 使用本组刚产生的抓取方案，避免泄漏旧答案。',
                   '检查 schema、候选资格、坐标引用链、提拉与释放顺序。',
                   '额外用当前图像的近似区域复核袖口、胸前目标与向内方向；未运行 IK 或物理动作。'],size=21)
    s=deck.slide('对照实验：总耗时与交互次数','A/B/I 包含两次调用；其余为一次完整选点与动作输出。',str(out/'analysis.json'))
    deck.table(s,['实验','秒','调用','响应轮次','Read / 看图','质量'],[[r['label'],f'{r["wall_s"]:.1f}',r['calls'],r['rounds'],f'{r["read_calls"]} / {r["view_calls"]}',r['quality']] for r in rows],widths=[4.3,1.2,.8,1.1,1.85,2.9],h=4.9,size=15)
    s=deck.slide('耗时比较：先观察效果，再讨论归因','条形长度是该组总墙钟时间，包含调用与本地校验。',str(out/'analysis.json'))
    mx=max(r['wall_s'] for r in rows)
    for i,r in enumerate(rows):
        y=1.82+i*.53;deck.text(s,r['label'],.65,y,4.15,.45,17)
        deck.rect(s,4.85,y+.03,6.6*r['wall_s']/mx,.3,TEAL if r['quality']=='离线检查通过' else ORANGE)
        deck.text(s,f'{r["wall_s"]:.1f}s',11.65,y,.98,.4,15)
    s=deck.slide('成对比较：每次只解释对应变化',f'联合调用对照采用 {"C′（C 有服务重试）" if control=="C_repeat" else "C"}；顺序、缓存和单样本波动仍限制因果结论。',str(out/'analysis.json'))
    pairs=[('全部上下文交付','A_files_split','B_inline_split'),('合并调用','B_inline_split',control),('明确步骤',control,'D_ordered'),('代码辅助','D_ordered','E_geometry'),('简洁输出',control,'F_concise'),('直接输入图片',control,'G_direct'),('同配置波动','C_inline_merged','C_repeat')]
    if 'H_low_effort' in by:pairs.append(('推理强度',control,'H_low_effort'))
    if 'I_matched_inline' in by:pairs.append(('已读信息匹配','A_files_split','I_matched_inline'))
    def qshort(r):return '通过' if r['quality']=='离线检查通过' else '复核'
    deck.table(s,['改变','之前 → 之后 / 秒','耗时减少','离线质量：前 → 后'],[[name,f'{by[a]["wall_s"]:.1f} → {by[b]["wall_s"]:.1f}',f'{100*(1-by[b]["wall_s"]/by[a]["wall_s"]):+.1f}%',qshort(by[a])+' → '+qshort(by[b])] for name,a,b in pairs],widths=[3.2,3.2,2,3.75],h=4.7,size=15)
    s=deck.slide('服务异常：需要单独归类的第四种因素','不能把服务端重试造成的长尾全部归因于结构、算法或 Agent。',str(out/'analysis.json'))
    rr=[[r['label'],r['api_retry_count'],str([e['status'] for e in r['api_retries']]) if r['api_retries'] else '无记录'] for r in rows]
    deck.table(s,['实验','API 重试次数','记录的错误状态'],rr,widths=[6.6,2.3,3.25],h=4.45,size=15)
    deck.text(s,'重试事件的时间戳不等于失败请求时长；不直接扣除 168 秒或将其视为纯网络时间。',.7,6.48,12,.4,14,ORANGE)
    s=deck.slide('Token 不是时间：这里展示 CLI 的记录值','主调用总量含缓存；thinking 已在 output 中。费用含辅助模型。* 表示 usage 可疑，不能推断推理量。',str(out/'analysis.json'))
    deck.table(s,['实验','主调用 token','输出 token','其中 thinking','记录费用 / USD'],[[code(r)+('*' if r['token_usage_caveats'] else ''),f'{r["tokens"]:,}',f'{r["output_tokens"]:,}',f'{r["reported_thinking_tokens"]:,}',f'{r["reported_cost_usd"]:.3f}'] for r in rows],widths=[1.2,3.1,2.4,2.5,2.95],h=4.9,size=15)
    s=deck.slide('等待、输出和传输：不能叫作纯思考时间','下表按观测窗口统计；不同列有包含关系，不用于相加得到总耗时。',str(out/'analysis.json'))
    deck.table(s,['实验','总墙钟 / 秒','响应前等待合计','响应输出窗口','CLI 窗口'],[[code(r),f'{r["wall_s"]:.1f}',f'{r["response_wait_s"]:.1f}',f'{r["emission_window_s"]:.1f}',f'{r["remote_cli_s"]:.1f}'] for r in rows],widths=[1.1,2.3,3.05,2.8,2.9],h=4.9,size=16)
    # Comparison images are geometric overlays of model outputs, not generative edits.
    snapshot=Path(read(out/'design.json')['frozen_snapshot']);base_image=Image.open(snapshot/'prepared_handoff/image_6.png').convert('RGB')
    assets=out/'presentation_assets';assets.mkdir(exist_ok=True)
    font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',23)
    for r in rows:
        im=base_image.copy();draw=ImageDraw.Draw(im)
        for pt,color,label in [(r['grasp_collar_up'],'#F59E0B','G'),(r['target_collar_up'],'#00D8CC','T')]:
            if pt:
                x,y=pt;draw.ellipse((x-12,y-12,x+12,y+12),outline=color,width=5);draw.text((x+16,y-13),label,fill=color,font=font)
        if r['grasp_collar_up'] and r['target_collar_up']:draw.line((*r['grasp_collar_up'],*r['target_collar_up']),fill='#00D8CC',width=3)
        im.save(assets/(r['arm']+'.png'))
    for chunk in [rows[i:i+4] for i in range(0,len(rows),4)]:
        s=deck.slide('输出质量：抓取点 G → 目标点 T','黄色为抓取点，青色为目标点；所有结果映射到同一张当前衣服图。',str(out/'analysis.json'))
        for i,r in enumerate(chunk):
            x=.6+(i%2)*6.35;y=1.68+(i//2)*2.58
            s.shapes.add_picture(str(assets/(r['arm']+'.png')),Inches(x+.65),Inches(y),width=Inches(3.75))
            deck.text(s,f'{code(r)} · {r["grasp_id"]} · {r["quality"]}',x,y+2.11,6,.35,15)
    if (out/'target_review.png').exists():
        s=deck.slide('两个“需复核”，原因完全不同','保留原始区域检查，不把近似边界当作物理真值；以下为图像复核。',str(out/'visual_review.json'))
        s.shapes.add_picture(str(out/'target_review.png'),Inches(.65),Inches(1.8),width=Inches(12))
        deck.text(s,'H：上胸近领口，距近似 ROI 边界约 2px；几何上合理，放置效果待验证。\nB：明显落到另一侧已折袖子区域；不能当作质量不变的加速。',.7,5.95,12,1,19)
    if (out/'excluded_attempts.json').exists():
        s=deck.slide('另一个结构性问题：CLI 启动时输入未就绪','失败尝试单独保留，不把“30.9 秒退出”当成加速。',str(out/'excluded_attempts.json'))
        deck.bullets(s,['I 首次启动报：3 秒内没有收到 stdin；未得到模型响应或 token 终值。',
                       '在测试传输中先完整缓存文本，再以 stream-json 启动 CLI。',
                       '重跑前核对：提示词和 schema 与失败时完全一致。',
                       '这修正的是实验输入交付；没有修改生产规划策略。失败记录仍可复查。'],size=21)
    s=deck.slide('能确认什么，尚不能确认什么','区分软件问题、方法收益与服务行为，避免过度归因。',str(out/'analysis.json'))
    deck.table(s,['层面','这次能验证','这次不能证明'],[['结构','减少 Read 和调用是否改善本场景耗时','其他场景也稳定加速'],['方法','明确步骤 / 几何辅助的输出和耗时变化','现有算法已最优，或耗时是理论下限'],['Agent','输出约束 / 取图方式的行为差异','全部差异来自模型纯推理或服务排队'],['质量','离线契约、引用链、语义区域一致性','真实抓取成功率与物理折叠精度']],widths=[1.4,5.4,5.35],h=3.9,size=18)
    if 'I_matched_inline' in by and 'H_low_effort' in by:
        def change(a,b):return f'{100*(1-by[b]["wall_s"]/by[a]["wall_s"]):+.1f}%'
        s=deck.slide('回答最初的三个问题','百分比表示耗时减少；负数表示变慢。各项只针对当前样本。',str(out/'analysis.json'))
        findings=[['结构 / 信息交付',f'A → I：{change("A_files_split","I_matched_inline")}；{by["I_matched_inline"]["quality"]}', '更接近隔离文件读取的多轮交互成本'],
                  ['方法 / 算法',f'D → E 减少 {change("D_ordered","E_geometry")}；相对 {code(by[control])}：{change(control,"E_geometry")}', '有局部收益，不等于优于基本联合规划'],
                  ['Agent 执行设置',f'{code(by[control])} → H：{change(control,"H_low_effort")}；{by["H_low_effort"]["quality"]}','同模型请求 low effort；不能外推稳定精度'],
                  ['服务异常',f'C：{c["wall_s"]:.1f}s / C′：{cr["wall_s"]:.1f}s','C 有 524 重试；这是必须单列的因素']]
        deck.table(s,['问题','实验结果','解释边界'],findings,widths=[2.5,5.0,4.65],h=4.15,size=17)
        deck.text(s,'三类原因可以并存；本实验不会把剩余时间自动归给“算法必需”或“Agent 天性”。',.7,6.2,12,.55,17,BLUE)
    s=deck.slide('建议：按验证结果逐步落地','本次实验没有修改生产推理策略、没有更新正式 skill。',str(out/'analysis.json'))
    deck.bullets(s,['先保留“原始计时 + 公共工具记录 + token”的统一观测口径。',
                   '优先复测最快且通过离线检查的配置；再扩展不同朝向与遮挡场景。',
                   '若收益大于同配置波动，再单独接入生产流程；保留原流程回退。',
                   '上线前追加静态编译 / IK 检查和受控机器人验证，不能直接由速度决定。'],size=21)
    s=deck.slide('复查入口与实验限制','所有数值来自本地记录；原始请求、输出、工具轨迹和 token 均保留。',str(out))
    deck.text(s,'原始真实轮次\nresults/two_stage_iteration_timing_20261005_172835/\n\n离线对照与统计\nresults/latency_ablation_20261005/\n\n复现实验\nscripts/benchmark_planner_latency.py\n\n报告生成\nscripts/report_planner_latency.py',.7,1.8,12,4.85,20)
    target=out/'planner_latency_study.pptx';deck.prs.save(target)
    with (out/'comparison.csv').open('w',encoding='utf-8-sig',newline='') as f:
        keys=['arm','label','status','wall_s','calls','rounds','read_calls','view_calls','response_wait_s','emission_window_s','remote_cli_s','tokens','all_models_reported_tokens','output_tokens','reported_thinking_tokens','reported_cost_usd','grasp_id','quality','api_retry_count']
        w=csv.DictWriter(f,keys,extrasaction='ignore');w.writeheader();w.writerows(rows)
    with (out/'original_timing.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(['阶段','耗时秒','占比']);w.writerows((n,round(v,3),round(100*v/total,2)) for n,v in times)
    lines=['# Planner 耗时对照报告','',f'原始真实轮次：{total:.3f} 秒。离线对照不包含前置拍照/编辑、机器人或 evaluation。','',
           '| 实验 | 秒 | 响应轮次 | Read | 看图 | 总 token | 离线质量 |','|---|---:|---:|---:|---:|---:|---|']
    lines += [f'| {r["label"]} | {r["wall_s"]:.1f} | {r["rounds"]} | {r["read_calls"]} | {r["view_calls"]} | {r["tokens"]} | {r["quality"]} |' for r in rows]
    lines+=['','一次筛查不能证明稳定加速或物理精度。区域评估是当前图像的近似人工式标注；baseline 答案不作为唯一真值。',
            'CLI 时间与 token 仅为观测记录；不推断隐含思考主题，不将等待等同于纯网络时间。']
    (out/'report.md').write_text('\n'.join(lines),encoding='utf-8')
    print(target)

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);p.add_argument('--original',type=Path,required=True);a=p.parse_args();build(a.output,a.original)
