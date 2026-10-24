#!/usr/bin/env python3
"""Offline, evidence-linked meeting deck; no robot/model imports or calls."""
from pathlib import Path
import json,csv,shutil,hashlib,html,zipfile
from datetime import datetime
import numpy as np
from PIL import Image,ImageDraw,ImageFont
from pptx import Presentation
from pptx.util import Inches
ROOT=Path(__file__).resolve().parents[1]; R=ROOT/'reports'; O=R/'unfolding_meeting_brief_20261021'
for d in ['slides','videos']: (O/d).mkdir(parents=True,exist_ok=True)
F='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'; W,H=1920,1080
sources={}; slides=[]
def track(p):
 p=Path(p).resolve();sources[str(p)]=hashlib.sha256(p.read_bytes()).hexdigest();return p
def read(p):return json.loads(track(p).read_text())
def f(n):return ImageFont.truetype(F,n)
def text(im,s,x,y,size=28,color='#334155',width=1780):
 d=ImageDraw.Draw(im);cy=y
 for para in str(s).split('\n'):
  line=''
  for c in para:
   if d.textlength(line+c,font=f(size))>width:
    d.text((x,cy),line,font=f(size),fill=color);cy+=size+12;line=''
   line+=c
  d.text((x,cy),line,font=f(size),fill=color);cy+=size+12
 if cy>1020:raise ValueError('text overflow: '+str(s))
 return cy
def pic(im,p,box):
 with Image.open(track(p)) as src:a=src.convert('RGB')
 a.thumbnail((box[2],box[3]),Image.Resampling.LANCZOS);im.paste(a,(int(box[0]+(box[2]-a.width)/2),int(box[1]+(box[3]-a.height)/2)))
def base(title,sub):
 im=Image.new('RGB',(W,H),'#f7f9fc');d=ImageDraw.Draw(im);d.rectangle((0,0,W,13),fill='#187a91')
 text(im,title,65,40,48,'#102b3b');text(im,sub,65,112,27,'#536575')
 text(im,f'UNFOLDING  /  2026.10.21                                        {len(slides)+1:02d}',65,1025,20,'#64748b') if False else None
 d.text((65,1035),f'UNFOLDING  /  2026.10.21     |     {len(slides)+1:02d}',font=f(20),fill='#64748b');return im
def save(im,title,notes='',video=None):
 p=O/'slides'/f'{len(slides)+1:02d}.png';im.save(p);slides.append(dict(title=title,image=str(p.relative_to(O)),notes=notes,video=video))
def table(im,headers,rows,y=225,widths=None,size=26):
 widths=widths or [1780/len(headers)]*len(headers);d=ImageDraw.Draw(im)
 for j,row in enumerate([headers]+rows):
  x=65; yy=y+j*65;d.rectangle((65,yy,1855,yy+62),fill='#dcebf0' if j==0 else ('#ffffff' if j%2 else '#edf2f6'))
  for val,w in zip(row,widths):
   text(im,str(val),x+12,yy+14,size,'#102b3b',width=w-20);x+=w
sets=[('memory','历史 memory',['trajectory_memory_full_episode_v1_collection_aligned'],'01_full_memory'),
 ('goal','Goal-only Agent',['goal_only_full_episode_v1_restart02','goal_only_full_episode_v1_continue01'],'02_goal_only_agent'),
 ('good','One-shot 成功候选',['one_shot_goal_full_episode_v1_restart01']+[f'one_shot_goal_full_episode_v1_continue0{i}' for i in range(1,5)],'03_one_shot_12round'),
 ('bad','One-shot 退化失败',['one_shot_goal_full_episode_v1_restart02'],'04_one_shot_regression')]
rows=[];totals=[]
def stats(v):return dict(n=len(v),mean=float(np.mean(v)),median=float(np.median(v)),p95=float(np.percentile(v,95))) if v else dict(n=0,mean=0,median=0,p95=0)
for key,label,runs,vid in sets:
 sums=dict(key=key,label=label,wall_s=0.,initialization_s=0.,initialization_model_s=0.,initialization_robot_s=0.,initialization_review_s=0.,planner_s=0.,gate_s=0.,host_s=0.,observation_s=0.,robot_s=0.,manipulation_stage_approx_s=0.,return_stage_approx_s=0.,round_other_s=0.,actions=0,rounds=0)
 lat=[];actlat=[];allsummary=[]
 for run in runs:
  root=R/run;s=read(root/'episode_summary.json');allsummary.append(s)
  sums['wall_s']+=s['episode_wall_s'];sums['actions']+=s['physical_actions']
  init=s.get('initialization',s.get('initialization_result',{}))
  if (root/'initialization/result.json').exists():init=read(root/'initialization/result.json')
  sums['initialization_s']+=s.get('initialization_wall_s',init.get('elapsed_s',0.))
  sums['initialization_model_s']+=s.get('initialization_planning_s',s.get('initialization_agent_latency_s',0.))
  sums['initialization_robot_s']+=init.get('robot_execution_s',0.)
  sums['initialization_review_s']+=s.get('initialization_visual_review',{}).get('wait_s',0.)
  for folder in sorted(root.glob('round_[0-9][0-9][0-9]')):
   if not (folder/'timing.json').exists():continue
   t=read(folder/'timing.json');plan=read(folder/'plan.json') if (folder/'plan.json').exists() else {};ex=read(folder/'execution.json')
   plan_s=t.get('planning_s',t.get('agent_wall_s',0.));robot=t.get('robot_execution_s',0.)
   r=dict(method=label,run=run,round=folder.name,decision=plan.get('decision','NO_PLAN'),execution=ex.get('status'),
    planner_s=plan_s,gate_s=t.get('gate_wall_s',0.),host_s=t.get('host_validation_s',t.get('host_planning_s',0.)),observation_s=t.get('observation_s',0.),robot_s=robot,round_wall_s=t.get('round_wall_s',0.),manipulation_stage_approx_s=0.,return_stage_approx_s=0.)
   if plan_s>0:lat.append(plan_s)
   if ex.get('status')=='EXECUTED' and plan_s>0:actlat.append(plan_s)
   if (folder/'events.jsonl').exists() and robot:
    ev=[json.loads(l) for l in track(folder/'events.jsonl').read_text().splitlines()]
    stage=next((e['t_s'] for e in ev if e.get('kind')=='stage' and e.get('phase')=='return_perception'),None)
    if stage is not None:r['manipulation_stage_approx_s']=min(stage,robot);r['return_stage_approx_s']=robot-r['manipulation_stage_approx_s']
   r['round_other_s']=r['round_wall_s']-sum(r[k] for k in ['planner_s','gate_s','host_s','observation_s','robot_s'])
   assert r['round_other_s']>-.2
   rows.append(r);sums['rounds']+=1
   for k in ['planner_s','gate_s','host_s','observation_s','robot_s','manipulation_stage_approx_s','return_stage_approx_s','round_other_s']:sums[k]+=r[k]
 sums['outside_round_other_s']=sums['wall_s']-sums['initialization_s']-sums['initialization_review_s']-sum(x['round_wall_s'] for x in rows if x['method']==label)
 sums['planning_stats']=stats(lat);sums['executed_action_planning_stats']=stats(actlat)
 sums['calendar_span_s']=(datetime.fromisoformat(allsummary[-1]['finished_at'])-datetime.fromisoformat(allsummary[0]['started_at'])).total_seconds()
 sums['between_segments_s']=sums['calendar_span_s']-sums['wall_s']
 assert abs(sums['wall_s']-sum(sums[k] for k in ['initialization_s','initialization_review_s','planner_s','gate_s','host_s','observation_s','robot_s','round_other_s','outside_round_other_s']))<.01
 totals.append(sums)
 shutil.copyfile(track(R/'meeting_materials_20261020/videos'/f'{vid}.mp4'),O/'videos'/f'{key}.mp4')
with (O/'per_round_timing.csv').open('w') as fcsv:
 w=csv.DictWriter(fcsv,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
(O/'timing_totals.json').write_text(json.dumps(totals,ensure_ascii=False,indent=2))
with (O/'timing_totals.csv').open('w') as fc:
 simple=[{k:v for k,v in t.items() if not isinstance(v,dict)} for t in totals];w=csv.DictWriter(fc,fieldnames=list(simple[0]));w.writeheader();w.writerows(simple)
# 1: graph schematic from frozen real graph
GROOT=R/'arm6_incremental_replay_20261019_224938';g=read(GROOT/'graph_used.json');tr=read(R/'arm6_graph_incremental_20261019_v3/compiled/trajectory.json')
im=base('01  Graph：把状态与人工动作连成图','真实 replay 快照：93 nodes / 111 edges / 202 reference images；下图为实际子图。')
for j,n in enumerate(['n000','n001','n002']):
 node=next(a for a in g['nodes'] if a['node_id']==n);p=tr['observations'][node['representative_observations'][0]]['rgb']['path'];x=65+j*610
 pic(im,p,(x,225,560,360));text(im,n+'  '+node['task_context'],x,605,27,width=550)
 if j<2:text(im,'→',x+555,375,55,'#187a91',width=65)
text(im,'e000：袖片内折                         e001：另一侧袖片内折',360,720,30,'#187a91')
text(im,'在线：当前 RGB → 检索参考 → GPT 确认节点 → 选边 / grounding → Host → 执行',65,815,34)
text(im,'节点记录可见布层关系；边记录已保存动作。此旧 Graph 含 folding 数据，不等同于后面的 139 条 unfolding memory。',65,900,26,width=1780)
save(im,'Graph 示意',str(GROOT/'graph_used.json'))
# 2 matching failure current + six candidates
attempt=read(GROOT/'step_0000/matching/matching_attempts.json')['attempts'][0]
im=base('02  看起来相近，仍无法确认同一节点','保存的 GPT 判定：NO_MATCH；检索候选相近 ≠ 可确认动作可迁移。不是人工标注的识别错误率。')
pic(im,GROOT/'step_0000/before.png',(65,235,650,440));text(im,'CURRENT',65,190,30,'#187a91')
text(im,'模型拒绝依据：\n图案露出比例、相邻素色布片覆盖、\n下方折边形状无法同时对齐。\n\n只展示原 Top-10 中前 6 张；\n最终 NO_MATCH 针对整组候选。',65,705,27,width=650)
ids=attempt['image_ids'][1:7]
for i,oid in enumerate(ids):
 x=760+(i%3)*365;y=215+(i//3)*370
 pic(im,tr['observations'][oid]['rgb']['path'],(x,y+40,350,260));node=next(n['node_id'] for n in g['nodes'] if oid in n['representative_observations']);text(im,f'{i+1}. {node} / {oid}',x,y,24,width=350)
text(im,'检索 0.158 s  +  GPT 复核 120.916 s；无动作执行。',760,960,26,'#a4552c',width=1090)
save(im,'相近参考与 NO_MATCH',str(GROOT/'step_0000/matching/matching_attempts.json'))
# 3 historical data evidence
im=base('03  历史经验：before → executed action → after','冻结 memory：139 accepted unfolding steps / 16 sessions；不要求 exact state matching。')
p=R/'trajectory_memory_full_episode_v1_collection_aligned/round_000/selected_experiences/contact_sheet.png';pic(im,p,(65,195,1240,730))
text(im,'真实 episode 使用情况\n\n7 次动作 / 12 次经验引用\n7 个 distinct cited steps\n139 个历史 step 被访问\n208 张去重源图被读取\n\n引用由原 plan.json 解析；\n读取不等于经验造成了改善。',1350,240,30,width=500)
save(im,'历史记录示例',str(p))
# 4-7 embedded movie slides
subs=[('历史 memory：7 步展开后人工停止','139 条经验 + 当前 RGB-D + 本 episode 历史','USER_STOP；最终人工 FOLD_READY 待判'),
('Goal-only Agent：直接规划','当前 RGB-D + 固定目标 + episode history；无历史 memory','2 + 6 动作；深度拒绝后经授权换视角续跑，最终 DONE'),
('One-shot：展开效果较好的成功候选','current RGB + goal RGB → 单次 direct GPT；12 轮，8 动作','含 4 次拒绝和 4 段人工续跑；终止 ROUND_LIMIT，人工待确认'),
('One-shot：覆盖退化的失败案例','同一 one-shot policy；round_003 为重点复盘动作','6 动作后安全停止；箭头动作把已展开区域重新覆盖')]
for index,((key,label,runs,vid),(title,sub,end)) in enumerate(zip(sets,subs)):
 im=base(f'{index+4:02d}  '+title,sub)
 poster=R/runs[0]/'round_000/before/rgb.png'
 pic(im,poster,(65,195,1290,726));t=totals[index]
 text(im,'点击视频播放\n\n'+end+f"\n\n规划均值 {t['planning_stats']['mean']:.2f} s\n运行累计 {t['wall_s']/60:.2f} min",1390,245,29,width=455)
 text(im,'保存的 OBS 组成的幻灯片视频；非连续运动录像，播放速度不代表真实耗时。',65,950,26)
 save(im,title,'Sources: '+', '.join(runs),video=f'videos/{key}.mp4')
#8 failure critical beforeafter
im=base('08  失败动作复盘：已展开区域被重新覆盖','One-shot restart02 / round_003：与最后的 Host 安全停止分开讨论。')
p=R/'one_shot_goal_full_episode_v1_restart02/round_003'
pic(im,p/'proposed_action_overlay.png',(65,230,870,580));pic(im,p/'after/rgb.png',(980,230,870,580))
text(im,'动作前 + 原始 Pa→Pb',65,185,30);text(im,'动作后',980,185,30)
text(im,'人工先前指出左侧展开区域被覆盖。展示原始前后图，不把 Host 拒绝或执行完成当作策略质量标签。',65,865,30)
save(im,'失败动作前后',str(p))
#9 latency overview
a=[t['planning_stats'] for t in totals]
im=base('09  耗时总览：主要差异在规划侧','同一模型系列 / high；不同初态、输入和干预，属于代表案例描述，不是严格消融。')
table(im,['方法','调用数','规划 median','规划 mean','规划 p95','运行累计'],[[t['label'],t['planning_stats']['n'],f"{t['planning_stats']['median']:.2f}s",f"{t['planning_stats']['mean']:.2f}s",f"{t['planning_stats']['p95']:.2f}s",f"{t['wall_s']/60:.2f}min"] for t in totals],widths=[440,180,285,285,285,305],size=28)
text(im,'仅已执行 ACTION 的规划均值：',65,640,32,'#187a91')
text(im,'  /  '.join(f"{t['label']} {t['executed_action_planning_stats']['mean']:.2f}s" for t in totals),65,700,28)
text(im,'Agent wall 包含启动、工具、读图与模型等待；内部模型 API 次数及纯推理时长无法从现有日志拆开。\nOne-shot wall 为单次 direct 请求；不含机器人时间。四个代表 run 均早于 shadow gate，gate=0。',65,815,28)
save(im,'规划 latency 汇总','per_round_timing.csv; timing_totals.json')
#10 time breakdown
im=base('10  总时长拆分：互斥项可加回运行总时长','单位：秒；续跑间人工等待不计入“运行累计”，单列于详细报告。')
labels=[('initialization_s','初始化（含其模型/机器人）'),('initialization_review_s','初始化确认等待'),('planner_s','展开规划 Agent / direct'),('host_s','Host 准备/几何检查'),('observation_s','每轮记录的采图'),('robot_s','操作 + 返回观测位'),('round_other_s','轮内其他 / 未单列'),('outside_round_other_s','轮外启动收尾 / 未单列'),('wall_s','运行总时长')]
table(im,['互斥时间项','Memory','Goal-only','One-shot 好','One-shot 失败'],[[label]+[f'{t[k]:.2f}' for t in totals] for k,label in labels],y=190,widths=[620,290,290,290,290],size=25)
text(im,'初始化模型是初始化总时间的子集；机器人返回是机器人总时间的子集，不能重复相加。',65,895,27)
text(im,'精细到逐轮、初始化子项、动作/返回阶段、跨段等待：见 timing_report.md 与 per_round_timing.csv。',65,950,26)
save(im,'互斥耗时拆分','timing_totals.csv')
#11 graph runtime and model limitations
im=base('11  Graph 与机器人阶段：再细一层','Graph 采用上面的同一 NO_MATCH run；机器人子阶段来自 recorder stage 时间戳，为近似分段。')
rt=attempt['retrieval_timing'];gt=read(GROOT/'timing.json')
table(im,['Graph 在线项目','秒'],[[k,f'{rt[k]:.6f}'] for k in ['query_image_decode_s','dino_preprocess_s','dino_forward_s','ann_search_s']]+[['检索总计（包含以上项）',f"{rt['retrieval_total_s']:.6f}"],['GPT rerank',f"{attempt['gpt_rerank_s']:.3f}"],['全 run（含启动/采图/选择）',f"{gt['total_s']:.3f}"]],y=210,widths=[730,240],size=25)
text(im,'机器人执行子阶段（s）',1100,220,30,'#187a91',width=740)
for i,t in enumerate(totals):
 text(im,t['label']+f"\n操作阶段 ≈{t['manipulation_stage_approx_s']:.1f}  / 返回阶段 ≈{t['return_stage_approx_s']:.1f}",1100,305+i*130,27,width=740)
text(im,'返回阶段含回 Home/观测位、松爪及少量收尾；不是纯运动控制器时长。Graph 离线构建不含在在线耗时中。',65,930,26)
save(im,'细分耗时',str(GROOT/'timing.json'))
# Written report with exact provenance
lines=['# 展开实验会议材料：耗时口径与原始证据','','单位均为秒。PPT 视频是 OBS 幻灯片，不能按播放长度衡量真实执行速度。',
'用户指定 12 轮 run 为成功候选；其终止为 ROUND_LIMIT，尚无人类 FOLD_READY 标签。',
'所有选择保留拒绝和人工续跑。未追加模型调用、未连接机器人。','','## 互斥总量','',
'|时间项|'+'|'.join(t['label'] for t in totals)+'|','|---|'+'---:|'*4]
for k,label in labels:lines.append('|'+label+'|'+'|'.join(f'{t[k]:.3f}' for t in totals)+'|')
lines+=['','初始化 / 规划 / Host / 采图 / 机器人 / 剩余时间互斥；gate 在所选 run 未启用。round_other 是 round wall 减已计项，outside_round_other 是整个 run 减初始化、确认和 round wall。含未完整计时的中止采图、启动/收尾等；不是纯 overhead 测量。',
'','## 规划与机器人子项','','|方法|n|median|mean|p95|已执行 ACTION 规划均值|操作阶段≈|返回阶段≈|','|---|---:|---:|---:|---:|---:|---:|---:|']
for t in totals:
 s=t['planning_stats'];lines.append(f"|{t['label']}|{s['n']}|{s['median']:.3f}|{s['mean']:.3f}|{s['p95']:.3f}|{t['executed_action_planning_stats']['mean']:.3f}|{t['manipulation_stage_approx_s']:.3f}|{t['return_stage_approx_s']:.3f}|")
lines+=['','机器人分段依据每轮 events.jsonl 的 return_perception stage 相对时间；recorder 与 robot timer 起点略不同，属于近似分段，返回差额含收尾。主表 robot_execution_s 使用原实测值。',
'Agent 的 shell/图片读取事件没有完整独立时钟；不能拆出纯 memory retrieval、纯 GPT thinking、网络等待，更不能据此把差值全部归因于 memory。',
'','## 初始化与跨段暂停（不重复相加）','']
for t in totals:
 lines += [f"- {t['label']}：初始化 {t['initialization_s']:.3f}；其中已记录模型 {t['initialization_model_s']:.3f}；"+(f"其中已记录机器人 {t['initialization_robot_s']:.3f}。" if t['initialization_robot_s'] else '初始化机器人未单独计时。')+f"另外初始化确认等待 {t['initialization_review_s']:.3f}。运行段合计 {t['wall_s']:.3f}；首段启动到末段结束 {t['calendar_span_s']:.3f}；段间暂停 {t['between_segments_s']:.3f}（未加入运行累计）。"]
lines+=['','Memory / goal-only 初始化 model 字段为原 summary 口径，初始化中的 comparison 等其他调用未完整独立汇总，不能当作所有初始化模型时间。',
'','## Graph 在线拆分','',json.dumps(dict(retrieval=rt,gpt_rerank_s=attempt['gpt_rerank_s'],matching_s=gt['steps'][0]['phases_s']['matching'],selection_s=gt['steps'][0]['phases_s']['selection'],startup_and_other_s=gt['total_s']-gt['steps'][0]['elapsed_s'],total_s=gt['total_s']),indent=2),
'','查询与参考图来自原 matching_attempts.json 指定的 image_ids，未重跑识别。NO_MATCH 是模型原判定，不等于经过真值验证的 false negative。Graph 的离线建库、编码、编译不是本次在线测量。',
'','## 来源','']
for key,label,runs,vid in sets:
 lines.append(f'### {label}');lines.extend(f'- [{run}](../{run}/report.md)' for run in runs)
lines+=['',f'- Graph 原记录：[{GROOT.name}](../{GROOT.name}/step_0000/matching/matching_attempts.json)',
'- [逐轮耗时 CSV](per_round_timing.csv) / [总量 CSV](timing_totals.csv) / [原始来源 SHA256](source_manifest.json)']
(O/'timing_report.md').write_text('\n'.join(lines),encoding='utf-8')
# PowerPoint with embedded mp4s, static PDF preview and portable web review
prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
for s in slides:
 slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.shapes.add_picture(str(O/s['image']),0,0,width=prs.slide_width,height=prs.slide_height)
 if s['video']:
  key=Path(s['video']).stem;run=next(a[2][0] for a in sets if a[0]==key);poster=R/run/'round_000/before/rgb.png'
  slide.shapes.add_movie(str(O/s['video']),Inches(65/120),Inches(195/120),Inches(1290/120),Inches(726/120),poster_frame_image=str(poster),mime_type='video/mp4')
 slide.notes_slide.notes_text_frame.text=s['notes']+'\nVideo is observation slideshow, not real-time footage. Timing: timing_report.md'
prs.save(O/'unfolding_brief.pptx')
images=[Image.open(O/s['image']).convert('RGB') for s in slides];images[0].save(O/'unfolding_brief_preview.pdf',save_all=True,append_images=images[1:],resolution=120.)
page=['<!doctype html><meta charset="utf-8"><title>Unfolding meeting brief</title><style>body{background:#10202c;color:white;font:18px sans-serif;margin:30px auto;max-width:1200px}img,video{width:100%}section{margin-bottom:48px}a{color:#8dd9e9}</style><h1>展开实验会议材料</h1><p><a href="unfolding_brief.pptx">PPT（内嵌视频）</a> · <a href="unfolding_brief_preview.pdf">PDF 预览</a> · <a href="timing_report.md">详细耗时</a></p>']
for s in slides:
 page+=['<section><h2>'+html.escape(s['title'])+'</h2><img src="'+s['image']+'">']
 if s['video']:page+=['<video controls preload="metadata" src="'+s['video']+'"></video>']
 page+=['</section>']
(O/'index.html').write_text('\n'.join(page),encoding='utf-8')
(O/'source_manifest.json').write_text(json.dumps(dict(sources=sources,slides=slides,selection=sets,model_calls=0,robot_connections=0),ensure_ascii=False,indent=2))
with zipfile.ZipFile(O/'unfolding_brief.pptx') as z:
 movies=[n for n in z.namelist() if n.endswith('.mp4')];assert len(movies)==4
assert all(hashlib.sha256(Path(p).read_bytes()).hexdigest()==h for p,h in sources.items())
# Compact montage for visual inspection
montage=Image.new('RGB',(1280,360*((len(images)+1)//2)),'#ddd')
for i,im in enumerate(images):im.thumbnail((640,360));montage.paste(im,((i%2)*640,(i//2)*360))
montage.save(O/'slide_overview.jpg')
print(json.dumps(dict(slides=len(slides),embedded_videos=len(movies),output=str(O),totals=totals),ensure_ascii=False,indent=2))
