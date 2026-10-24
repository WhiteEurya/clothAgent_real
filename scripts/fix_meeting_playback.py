#!/usr/bin/env python3
"""Offline playback repair: autoplay video variant and video-free native slide sequence."""
from pathlib import Path
import json,re,zipfile
from copy import deepcopy
from lxml import etree
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.dml.color import RGBColor
from pptx.util import Inches
ROOT=Path(__file__).resolve().parents[1];O=ROOT/'reports/unfolding_meeting_brief_20261021'
src=O/'unfolding_brief_editable_simple_videos.pptx';videos=json.loads((O/'videos_simple/manifest.json').read_text())['videos']
# Fix the pre-existing indefinite start on every actual video timing node.
p=Presentation(src);fixed=0
for s in p.slides:
 for cond in s._element.xpath('.//p:video/p:cMediaNode/p:cTn/p:stCondLst/p:cond'):
  assert cond.get('delay')=='indefinite';cond.set('delay','0');fixed+=1
assert fixed==4
auto=O/'unfolding_brief_video_autoplay.pptx';p.save(auto)
# Expand only movie pages into actual photo slides. No movie decoder or poster involved.
p=Presentation(src);ids=p.slides._sldIdLst;old=list(ids);order=[];sequences=[];movie_index=0
for sid in old:
 s=p.part.related_slide(sid.rId)
 if not any(sh._element.xpath('.//a:videoFile') for sh in s.shapes):order.append(sid);continue
 item=videos[movie_index];movie_index+=1;start=len(order)+1
 for i,frame in enumerate(item['frames']):
  ns=p.slides.add_slide(p.slide_layouts[6]);ns.background.fill.solid();ns.background.fill.fore_color.rgb=RGBColor(0,0,0)
  # Exactly the same uncropped normalized image used to encode the MP4.
  image=O/'videos_simple'/item['key']/f'{i:03}.png'
  ns.shapes.add_picture(str(image),0,0,width=p.slide_width,height=p.slide_height)
  transition=OxmlElement('p:transition');transition.set('advClick','1');transition.set('spd','fast')
  if i<len(item['frames'])-1:transition.set('advTm','1200')
  ns._element.append(transition)
  ns.notes_slide.notes_text_frame.text=json.dumps(dict(method=item['label'],sequence_frame=i+1,frame_count=len(item['frames']),source=frame['source'],source_kind=frame['kind'],auto_advance_ms=1200 if i<len(item['frames'])-1 else None,last_frame='Pause for presenter; press arrow to continue',manual_continuations='Recorded original segments; not an uninterrupted autonomous episode'),ensure_ascii=False,indent=2)
  order.append(list(ids)[-1])
 sequences.append(dict(key=item['key'],first_slide=start,last_slide=len(order),frames=len(item['frames'])))
 p.part.drop_rel(sid.rId)
for sid in list(ids):ids.remove(sid)
for sid in order:ids.append(sid)
for i,s in enumerate(p.slides):
 for sh in s.shapes:
  if sh.has_text_frame and sh.text:
   if sh.top==Inches(35/120):sh.text_frame.paragraphs[0].runs[0].text=f'{i+1:02d}  '+re.sub(r'^\d{2}  ','',sh.text)
   elif sh.text.startswith('UNFOLDING  /'):sh.text_frame.paragraphs[0].runs[0].text=f'UNFOLDING  /  2026.10.21    |    {i+1:02d}    |    可编辑版'
# Explicitly honor stored timings in standard presentation properties.
for rel in p.part.rels.values():
 if rel.reltype.endswith('/presProps'):
  part=rel.target_part;xml=etree.fromstring(part.blob);tag='{http://schemas.openxmlformats.org/presentationml/2006/main}showPr'
  sp=xml.find(tag)
  if sp is None:
   sp=etree.Element(tag);ext=xml.find('{http://schemas.openxmlformats.org/presentationml/2006/main}extLst')
   if ext is None:xml.append(sp)
   else:xml.insert(list(xml).index(ext),sp)
  sp.set('useTimings','1');part._blob=etree.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
photo=O/'unfolding_brief_native_photo_playback.pptx';p.save(photo)
check=Presentation(photo);assert len(check.slides)==57
assert sum(sh.has_table for s in check.slides for sh in s.shapes)==3
photo_pages=0;auto_pages=0
for s in check.slides:
 assert not s._element.xpath('.//a:videoFile')
 ts=s._element.xpath('./p:transition')
 if ts:
  photo_pages+=1;auto_pages+=int(ts[0].get('advTm')=='1200')
  assert len(s.shapes)==1 and s.shapes[0].shape_type==13
assert photo_pages==42 and auto_pages==38
with zipfile.ZipFile(photo) as z:
 assert not any(n.endswith('.mp4') for n in z.namelist());assert z.testzip() is None
 with zipfile.ZipFile(auto) as a:assert sum(n.endswith('.mp4') for n in a.namelist())==4
validation=dict(video_autoplay_file=str(auto),native_photo_file=str(photo),slides=57,photo_frames=42,automatic_advances=38,manual_last_frames=4,native_tables=3,sequences=sequences,embedded_videos_in_photo_version=0,autoplay_conditions_fixed=4,validation='Package structure and image sequencing checked; target WPS slideshow playback not yet interactively verified',timing_reference='https://learn.microsoft.com/en-us/dotnet/api/documentformat.openxml.presentation.transition?view=openxml-3.0.1')
(O/'playback_fix_verification.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2))
index=O/'index.html';html=index.read_text().replace('最新版：纯图片短视频 + 团叠初态（可编辑 PPT）','旧版：短视频（未设自动播放）')
html=html.replace('<h1>展开实验会议材料</h1>','<h1>展开实验会议材料</h1><p><a href="unfolding_brief_native_photo_playback.pptx">最新版：原生图片顺序翻页（无嵌入视频）</a> · <a href="unfolding_brief_video_autoplay.pptx">视频自动启动修正版</a></p>');index.write_text(html)
print(json.dumps(validation,ensure_ascii=False,indent=2))
