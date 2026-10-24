#!/usr/bin/env python3
"""Replace the four meeting videos with short, unannotated observation slideshows."""
import hashlib,json,subprocess,zipfile
from pathlib import Path
from PIL import Image,ImageOps
from pptx import Presentation
from pptx.util import Inches
ROOT=Path(__file__).resolve().parents[1];R=ROOT/'reports';O=R/'unfolding_meeting_brief_20261021'
V=O/'videos_simple';V.mkdir(exist_ok=True)
selection=json.loads((O/'source_manifest.json').read_text())['selection'];items=[]
for key,label,runs,_ in selection:
 frames=[]
 def add(p,kind):
  if p.exists():frames.append(dict(source=str(p.resolve()),kind=kind,sha256=hashlib.sha256(p.read_bytes()).hexdigest()))
 add(R/runs[0]/'initialization/before/rgb.png','before_initialization')
 for run in runs:
  root=R/run; rounds=sorted(root.glob('round_[0-9][0-9][0-9]'))
  if rounds:add(rounds[0]/'before/rgb.png','segment_first_observation')
  for rd in rounds:
   ep=rd/'execution.json'
   if ep.exists() and json.loads(ep.read_text()).get('status')=='EXECUTED':add(rd/'after/rgb.png','after_executed_action')
 work=V/key;work.mkdir(exist_ok=True);lines=[]
 for i,frame in enumerate(frames):
  with Image.open(frame['source']) as im:
   # Full uncropped observation, no drawn text/arrow, no title or summary frame.
   im=ImageOps.pad(im.convert('RGB'),(1280,720),color='black');im.save(work/f'{i:03}.png')
  lines += [f"file '{i:03}.png'",'duration 1.2']
 lines += [f"file '{len(frames)-1:03}.png'"]
 (work/'frames.txt').write_text('\n'.join(lines)+'\n');out=V/f'{key}.mp4';duration=len(frames)*1.2
 subprocess.run(['ffmpeg','-hide_banner','-loglevel','error','-y','-f','concat','-safe','0','-i',str(work/'frames.txt'),'-t',str(duration),'-r','25','-c:v','libx264','-preset','fast','-crf','19','-pix_fmt','yuv420p','-movflags','+faststart','-an',str(out)],check=True)
 probe=json.loads(subprocess.check_output(['ffprobe','-v','error','-show_entries','format=duration:stream=width,height,codec_type','-of','json',str(out)]))
 assert abs(float(probe['format']['duration'])-duration)<.05
 assert len(probe['streams'])==1 and probe['streams'][0]['codec_type']=='video'
 items.append(dict(key=key,label=label,video=str(out),seconds_per_image=1.2,duration_s=duration,frame_count=len(frames),frames=frames))
 print(key,len(frames),duration,flush=True)
src=O/'unfolding_brief_editable_crumpled_initials.pptx';prs=Presentation(src);replacements={};found=[]
for s in prs.slides:
 movies=[sh for sh in s.shapes if sh._element.xpath('.//a:videoFile')]
 if not movies:continue
 assert len(movies)==1
 item=items[len(found)];sh=movies[0];node=sh._element.xpath('.//a:videoFile')[0]
 rid=node.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}link')
 part=s.part.rels[rid].target_part
 replacements[str(part.partname).lstrip('/')]=Path(item['video']).read_bytes()
 # Keep editable title/footer; remove the long sidebar and explanatory body.
 for other in list(s.shapes):
  if not other.has_text_frame:continue
  if other.top==Inches(35/120) or other.text.startswith('UNFOLDING  /'):continue
  if other.top==Inches(115/120):
   other.text_frame.paragraphs[0].runs[0].text=f"观测图顺序轮播 · 每张 1.2 秒 · 共 {item['duration_s']:.1f} 秒"
   for paragraph in list(other.text_frame.paragraphs)[1:]:paragraph._p.getparent().remove(paragraph._p)
  else:s.shapes._spTree.remove(other._element)
 sh.left=Inches(245/120);sh.top=Inches(195/120);sh.width=Inches(1430/120);sh.height=Inches(804.375/120)
 s.notes_slide.notes_text_frame.text += '\nVideo simplified: only original observation images; 1.2 s each; no burned-in labels, diagrams or narration. Accelerated slideshow, not real-time footage. Source segments and manual continuations: videos_simple/manifest.json. Actual timing remains in dedicated report slides.'
 found.append(item['key'])
assert found==[x['key'] for x in items]
tmp=V/'deck_before_media_replacement.pptx';prs.save(tmp)
out=O/'unfolding_brief_editable_simple_videos.pptx'
with zipfile.ZipFile(tmp) as srczip,zipfile.ZipFile(out,'w',compression=zipfile.ZIP_DEFLATED) as target:
 for info in srczip.infolist():target.writestr(info,replacements.get(info.filename,srczip.read(info.filename)))
with zipfile.ZipFile(out) as z:
 assert z.testzip() is None
 assert len([p for p in z.namelist() if p.endswith('.mp4')])==4
 for name,data in replacements.items():assert z.read(name)==data
check=Presentation(out);assert len(check.slides)==19
assert sum(sh.has_table for s in check.slides for sh in s.shapes)==3
(V/'manifest.json').write_text(json.dumps(dict(presentation=str(out),rule='initialization-before, first observation per saved segment, then after each EXECUTED action; chronological order within original segments; no annotated frames',videos=items),ensure_ascii=False,indent=2))
index=O/'index.html';html=index.read_text();html=html.replace('最新版：仅使用团叠初态的可编辑 PPT（19 页）','旧版：团叠初态版（视频已在新版简化）')
html=html.replace('<h1>展开实验会议材料</h1>','<h1>展开实验会议材料</h1><p><a href="unfolding_brief_editable_simple_videos.pptx">最新版：纯图片短视频 + 团叠初态（可编辑 PPT）</a></p>')
for item in items:html=html.replace(f'src="videos/{item["key"]}.mp4"',f'src="videos_simple/{item["key"]}.mp4"')
index.write_text(html)
print(out)
