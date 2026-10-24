#!/usr/bin/env python3
"""Add editable, source-traceable visual state evidence to the meeting deck. Offline only."""
import json, runpy, hashlib, zipfile
from pathlib import Path
from PIL import Image
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1]
ns=runpy.run_path(str(ROOT/'scripts/build_unfolding_meeting_editable.py'))
prs=ns['prs']; O=ns['O']; R=ns['R']
base,text,rect,picture,pos=[ns[k] for k in ['base','text','rect','picture','pos']]
read=ns['read']; evidence=[]
def source(p,**extra):
 p=Path(p).resolve();return dict(path=str(p),sha256=hashlib.sha256(p.read_bytes()).hexdigest(),**extra)
def cropped(s,p,box,crop=(200,40,1100,720)):
 # Native PPT picture crop: source pixels remain embedded and independently editable.
 iw,ih=Image.open(p).size; l,t,r,b=crop
 x,y,w,h=box;k=min(w/(r-l),h/(b-t));ww,hh=(r-l)*k,(b-t)*k
 sh=s.shapes.add_picture(str(p),pos(x+(w-ww)/2),pos(y+(h-hh)/2),width=pos(ww),height=pos(hh))
 sh.crop_left=l/iw;sh.crop_right=(iw-r)/iw;sh.crop_top=t/ih;sh.crop_bottom=(ih-b)/ih
 return sh
starts=[
 ('S01','One-shot · 10/20 22:02','one_shot_goal_full_episode_v1_20261020_220217_032210193'),
 ('S02','One-shot · restart01','one_shot_goal_full_episode_v1_restart01'),
 ('S03','One-shot · restart02','one_shot_goal_full_episode_v1_restart02'),
 ('S04','One-shot · 10/20 22:53','one_shot_goal_full_episode_v1_20261020_225322_882924401'),
 ('S05','Continuous · 11:30 / ep02','one_shot_goal_continuous_v1_20261021_113055_912326188/episode_0002'),
 ('S06','Continuous · 11:30 / ep01','one_shot_goal_continuous_v1_20261021_113055_912326188/episode_0001'),
 ('S07','Continuous · 10:23 / ep01','one_shot_goal_continuous_v1_20261021_102354_200485125/episode_0001'),
 ('S08','Continuous · 00:05 / ep01','one_shot_goal_continuous_v1_20261021_000504_812775375/episode_0001'),
 ('S09','Continuous · 11:05 / ep01','one_shot_goal_continuous_v1_20261021_110536_963158048/episode_0001'),
 ('S10','Memory · collection aligned','trajectory_memory_full_episode_v1_collection_aligned'),
 ('S11','Goal-only · restart02','goal_only_full_episode_v1_restart02'),
]
for page in range(2):
 batch=starts[page*6:page*6+6]
 s=base('展开前的团叠初态：11 次真实初始化 '+str(page+1)+'/2',
 '抓起再放下后的第一帧；尚未执行展开动作。只展示团叠初态。',notes='Selected real episode starts; not a random sample, not a count of all possible states. All 11 sources have saved EXECUTED initialization and are round_000/before, before any unfolding action. Native equal crop (220,70,1080,720), no rotation or generated images.')
 for j,(sid,label,run) in enumerate(batch):
  root=R/run;p=root/'round_000/before/rgb.png';x=65+j%3*605;y=205+j//3*390
  rect(s,x,y,580,370,'#ffffff','#d5e0e6',True)
  text(s,sid+'  '+label,x+16,y+12,550,42,24,'#187a91',True)
  cropped(s,p,(x+12,y+58,556,298),crop=(220,70,1080,720))
  summary=read(root/'episode_summary.json');init=root/'initialization/result.json'
  assert read(init)['status']=='EXECUTED'
  assert summary['initialization_physical_actions']>=1
  evidence.append(dict(kind='episode_start',id=sid,image=source(p),episode_summary=source(root/'episode_summary.json'),initialization=source(init) if init.exists() else None,caption=label))
 text(s,'11 张不同 episode 的初始团叠图；配对页复用这些图，不把重复展示计为新样本。',65,1000,1790,33,22,'#64748b')
lookup={sid:(R/run/'round_000/before/rgb.png') for sid,label,run in starts}
pairs=[
 ('都是紧凑布团：表面的折边不同','S01','S04','斜向宽布片盖在表层','环状折边围出中央凹口','看表层折边：斜盖片 / 环状凹口。'),
 ('都是多角团叠：外伸布片不同','S02','S08','下方宽布片向外突出','两侧布片伸出、下方较窄','看外伸片：位置与宽度不同。'),
 ('都是印花局部露出：前景覆盖不同','S03','S06','两侧高折边围住印花','印花下方横着一条长折片','看印花下方：包围式折边 / 横长折片。'),
 ('都是上团下片：底部形态不同','S05','S06','下方宽片呈扇形铺出','下方布片呈窄长弯折','看下方：宽扇形片 / 窄长弯折。'),
 ('都是紧凑印花团：可见面不同','S07','S09','宽素色面朝上，印花在侧边','印花面朝上，右侧团折','看表层：素色片 / 印花片。'),
 ('都是印花被遮：覆盖边方向不同','S10','S11','印花靠右，素色片在左下','印花靠下，上方宽片斜盖','看印花与遮挡布片的相对位置。'),
]
for page in range(3):
 s=base('同是团叠初态，局部关系不同 '+str(page+1)+'/3',
 '全部来自前两页的初始化后第一帧；两张图都尚未开始展开。',notes='Only crumpled initial observations S01–S11, after saved initialization and before any unfolding action. Pairs are selected visual illustrations, not matched-state ground truth. Captions describe visible cloth edges only; hidden topology and action transfer are not proven. See crumpled_initial_manifest.json for all source paths and hashes.')
 for j,(title,oa,ob,ca,cb,note) in enumerate(pairs[page*2:page*2+2]):
  y=205+j*395
  text(s,title,65,y,1790,48,32,'#102b3b',True)
  pa,pb=lookup[oa],lookup[ob]
  for x,p,oid,cap in [(65,pa,oa,ca),(620,pb,ob,cb)]:
   cropped(s,p,(x,y+65,520,290),crop=(220,70,1080,720));text(s,oid+'  '+cap,x,y+355,535,37,24,'#187a91')
  rect(s,1200,y+82,645,245,'#e5eef4',round=True)
  text(s,note,1225,y+115,595,110,29)
  text(s,'初始化后 · 展开前',1225,y+250,595,50,24,'#187a91',True)
  evidence.append(dict(kind='crumpled_initial_pair',title=title,source_ids=[oa,ob],images=[source(pa),source(pb)],caption=note,label_type='Visible RGB differences; not hidden-topology or action-transfer ground truth'))
 text(s,'整体都是“一团”，但露出的布片和覆盖边不同；RGB 不能证明隐藏层数或动作一定不可迁移。',65,1000,1790,33,22,'#64748b')
# New evidence pages follow the Graph overview, before the existing single-query evidence.
ids=prs.slides._sldIdLst;old=list(ids)
new=[old[i] for i in [0,1,14,15,16,17,18,2,3,4,5,6,7,8,9,10,11,12,13]]
for node in list(ids):ids.remove(node)
for node in new:ids.append(node)
import re
for i,s in enumerate(prs.slides):
 for sh in s.shapes:
  if not sh.has_text_frame:continue
  if sh.text.startswith('UNFOLDING  /'):sh.text_frame.paragraphs[0].runs[0].text=f'UNFOLDING  /  2026.10.21    |    {i+1:02d}    |    可编辑版'
  elif sh.top==pos(35):sh.text_frame.paragraphs[0].runs[0].text=f'{i+1:02d}  '+re.sub(r'^\d{2}  ','',sh.text)
out=O/'unfolding_brief_editable_crumpled_initials.pptx';prs.save(out)
manifest=dict(presentation=str(out),added_slides=5,episode_start_images=11,initialized_episodes=11,uninitialized_flat_comparison=0,crumpled_initial_pairs=6,added_image_placements=23,limitations=['Illustrative selection, not distribution or failure-rate estimate','Visible differences do not prove hidden topology or action non-transferability'],evidence=evidence)
(O/'crumpled_initial_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
check=Presentation(out)
assert len(check.slides)==19
assert sum(sh.has_table for s in check.slides for sh in s.shapes)==3
with zipfile.ZipFile(out) as z:
 assert z.testzip() is None
 movies=[n for n in z.namelist() if n.endswith('.mp4')];assert len(movies)==4
 with zipfile.ZipFile(O/'unfolding_brief_editable_sections.pptx') as original:
  assert sorted(hashlib.sha256(z.read(n)).hexdigest() for n in movies)==sorted(hashlib.sha256(original.read(n)).hexdigest() for n in original.namelist() if n.endswith('.mp4'))
for s in check.slides:
 assert sum(sh.has_text_frame for sh in s.shapes)>=4
 for sh in s.shapes:
  if sh.shape_type==13:assert sh.width<prs.slide_width*.95 or sh.height<prs.slide_height*.95
(O/'crumpled_initial_verification.json').write_text(json.dumps(dict(slides=19,added_evidence_pages=5,embedded_videos=4,video_bytes_unchanged=True,native_tables=3,native_text_and_picture_crops=True),indent=2))
print(out)
