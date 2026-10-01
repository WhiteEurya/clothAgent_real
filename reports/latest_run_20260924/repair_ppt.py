from pathlib import Path
from io import BytesIO
import zipfile,posixpath,xml.etree.ElementTree as E
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
out=Path(__file__).parent
p='http://schemas.openxmlformats.org/presentationml/2006/main';a='http://schemas.openxmlformats.org/drawingml/2006/main';r='http://schemas.openxmlformats.org/officeDocument/2006/relationships';ns={'p':p,'a':a}
prs=Presentation();prs.slide_width=12192000;prs.slide_height=6858000
with zipfile.ZipFile(out/'latest_run_report_with_process.pptx') as z:
 for i in range(1,10):
  slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(255,255,255)
  rels={x.get('Id'):posixpath.normpath('ppt/slides/'+x.get('Target')) for x in E.fromstring(z.read(f'ppt/slides/_rels/slide{i}.xml.rels'))}
  tree=E.fromstring(z.read(f'ppt/slides/slide{i}.xml'))
  for el in tree.find('p:cSld/p:spTree',ns):
   if el.tag not in [f'{{{p}}}sp',f'{{{p}}}pic']:continue
   xf=el.find('p:spPr/a:xfrm',ns);off=xf.find('a:off',ns);ext=xf.find('a:ext',ns);x,y,cx,cy=[int(v) for v in (off.get('x'),off.get('y'),ext.get('cx'),ext.get('cy'))]
   if el.tag==f'{{{p}}}pic':
    rid=el.find('p:blipFill/a:blip',ns).get(f'{{{r}}}embed');blob=z.read(rels[rid]);vid=el.find('.//a:videoFile',ns)
    if vid is not None:
     movie=out/'full_run_32x.mp4';slide.shapes.add_movie(str(movie),x,y,cx,cy,poster_frame_image=BytesIO(blob),mime_type='video/mp4')
    else:slide.shapes.add_picture(BytesIO(blob),x,y,cx,cy)
   else:
    border=el.find('p:spPr/a:ln/a:solidFill',ns) is not None
    shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,x,y,cx,cy) if border else slide.shapes.add_textbox(x,y,cx,cy)
    shape.fill.background()
    if border:shape.line.color.rgb=RGBColor(0,0,0);shape.line.width=Pt(1)
    else:shape.line.fill.background()
    tf=shape.text_frame;tf.word_wrap=True
    for j,para in enumerate(el.findall('p:txBody/a:p',ns)):
     pp=tf.paragraphs[0] if j==0 else tf.add_paragraph();pp.space_after=Pt(0)
     for rr in para.findall('a:r',ns):
      run=pp.add_run();run.text=rr.findtext('a:t','',ns);rp=rr.find('a:rPr',ns);run.font.size=Pt(int(rp.get('sz','2000'))/100);run.font.name='Microsoft YaHei';run.font.color.rgb=RGBColor(0,0,0)
prs.save(out/'latest_run_report_fixed.pptx')
check=Presentation(out/'latest_run_report_fixed.pptx');print('Reopened:',len(check.slides),'slides')
