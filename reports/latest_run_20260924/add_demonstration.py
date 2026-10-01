from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from PIL import Image
out=Path(__file__).parent;prs=Presentation(out/'latest_run_report_fixed.pptx')
def tx(slide,text,x,y,w,h,size):
 sh=slide.shapes.add_textbox(int(x*6350),int(y*6350),int(w*6350),int(h*6350));tf=sh.text_frame;tf.word_wrap=True
 for i,line in enumerate(text.split('\n')):
  pp=tf.paragraphs[0] if i==0 else tf.add_paragraph();pp.text=line
  for rr in pp.runs:rr.font.name='Microsoft YaHei';rr.font.size=Pt(size);rr.font.color.rgb=RGBColor(0,0,0)
# Preserve formatting while clarifying the story and algorithm.
for sl in prs.slides:
 for sh in sl.shapes:
  if not sh.has_text_frame:continue
  for pp in sh.text_frame.paragraphs:
   for rr in pp.runs:
    rr.text=rr.text.replace('预训练模型 + 真实反馈 + 执行约束 → 有效的 harness policy。','预训练模型 + demonstration + 真实反馈 → harness policy。').replace('当前阶段与历史经验','当前阶段、示范图与经验')
s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(255,255,255)
tx(s,'Demonstration：用示范图说明目标状态',65,45,1800,90,32)
tx(s,'每个折叠阶段提供“示范前 → 示范后”图像，帮助 Claude 理解该步目标。\n示范用于推理时参考，不进行任务策略训练，也不直接回放示范动作。\n实际抓点、深度和轨迹来自当前衣物观测；下图为本 run 使用的下摆示范。',70,145,1780,210,22)
r=Path('/mnt/newssd/sja/clothAgent_real/runs/2026-09-24/fold_20260924T014204289928777Z/results/fold_exploration/20260924T031106388377Z/iteration_008/fold_state_reference')
for j,(name,label) in enumerate([('fold_reference_source.png','示范前：完成两侧折叠'),('fold_reference_target.png','示范后：下摆对折')]):
 tx(s,label,200+j*870,365,820,55,21)
 path=r/name;iw,ih=Image.open(path).size;scale=min(820/iw,580/ih);w=iw*scale;h=ih*scale
 s.shapes.add_picture(str(path),int((200+j*870+(820-w)/2)*6350),int(430*6350),int(w*6350),int(h*6350))
# Insert immediately after the story; all media relationships remain intact.
ids=prs.slides._sldIdLst;node=ids[-1];ids.remove(node);ids.insert(1,node)
for i,sl in enumerate(prs.slides,1):
 for sh in sl.shapes:
  if sh.has_text_frame and sh.text.startswith('ClothAgent |'):
   for pp in sh.text_frame.paragraphs:
    for rr in pp.runs:rr.text=f'ClothAgent | 2026-09-24 run | {i}/10'
tx(s,'ClothAgent | 2026-09-24 run | 2/10',65,1020,1700,40,12)
prs.save(out/'latest_run_report_demonstration.pptx');print('Saved 10 slides with demonstration; video retained')
